import json
import math
import re
from collections import defaultdict
from pathlib import Path
from statistics import fmean

BASE_DIR = Path(__file__).resolve().parent
PRED_FOLDER = BASE_DIR / 'data' / 'pred'
GT_FOLDER = BASE_DIR / 'data' / 'gt'
OUTPUT_FOLDER = BASE_DIR / 'metrics_output'
# Same video and numeric frame ID match even with different zero padding/case.
# Optional suffixes: Video1_0001_pred.json / Video1_1_GT.json.
FILENAME_PATTERN = r'^(?P<video>.+?)[_-](?P<frame>\d+)(?:[_-](?:pred|gt))?$'
ALLOW_UNMATCHED = False
OVERWRITE = False

JOINTS = ('LWrist', 'RWrist', 'LElbow', 'RElbow', 'LShoulder', 'RShoulder',
          'LHip', 'RHip', 'LKnee', 'RKnee', 'LAnkle', 'RAnkle')
# End point, vertex, end point. Eight angles, including both elbows.
ANGLES = {
    'LElbow': ('LShoulder', 'LElbow', 'LWrist'),
    'LHip': ('LShoulder', 'LHip', 'LKnee'),
    'LKnee': ('LHip', 'LKnee', 'LAnkle'),
    'LShoulder': ('LElbow', 'LShoulder', 'LHip'),
    'RElbow': ('RShoulder', 'RElbow', 'RWrist'),
    'RHip': ('RShoulder', 'RHip', 'RKnee'),
    'RKnee': ('RHip', 'RKnee', 'RAnkle'),
    'RShoulder': ('RElbow', 'RShoulder', 'RHip'),
}
SUMMARY_HEADERS = ['video', 'frame_count', 'MPJPE_mean', 'MPJPE_min',
                   'MPJPE_max', 'MPJAE_mean', 'MPJAE_min', 'MPJAE_max']


def natural_key(value):
    return tuple((0, int(part)) if part.isdigit() else (1, part.casefold())
                 for part in re.split(r'(\d+)', str(value)))


def read_points(path):
    with Path(path).open(encoding='utf-8-sig') as handle:
        value = json.load(handle)
    validate_points(value, str(path))
    return {joint: [float(v) for v in value[joint]] for joint in JOINTS}


def validate_points(points, source='input'):
    if not isinstance(points, dict):
        raise ValueError(f'{source}: expected a landmark object')
    for joint in JOINTS:
        point = points.get(joint)
        if not isinstance(point, (list, tuple)) or len(point) != 3:
            raise ValueError(f'{source}: {joint} must contain exactly three coordinates')
        if any(isinstance(v, bool) or not isinstance(v, (int, float)) or
               not math.isfinite(v) for v in point):
            raise ValueError(f'{source}: {joint} contains non-finite/non-numeric coordinates')


def joint_angle(points, triplet):
    a, vertex, b = (points[name] for name in triplet)
    u = [x-y for x, y in zip(a, vertex)]
    v = [x-y for x, y in zip(b, vertex)]
    denominator = math.sqrt(sum(x*x for x in u)) * math.sqrt(sum(x*x for x in v))
    if denominator == 0:
        raise ValueError(f'Zero-length segment for angle {triplet}')
    cosine = sum(x*y for x, y in zip(u, v)) / denominator
    return math.degrees(math.acos(max(-1.0, min(1.0, cosine))))


def calculate_frame(pred, gt):
    validate_points(pred, 'pred')
    validate_points(gt, 'GT')
    result = {}
    for joint in JOINTS:
        for axis, i in zip('XYZ', range(3)):
            result[f'{joint}_{axis}_pred'] = pred[joint][i]
            result[f'{joint}_{axis}_GT'] = gt[joint][i]
    errors = []
    for name, triplet in ANGLES.items():
        pred_angle = joint_angle(pred, triplet)
        gt_angle = joint_angle(gt, triplet)
        error = abs(pred_angle-gt_angle)
        result.update({f'{name}_angle_pred': pred_angle,
                       f'{name}_angle_GT': gt_angle, f'{name}_angle_error': error})
        errors.append(error)
    result['MPJPE'] = fmean(math.dist(pred[joint], gt[joint]) for joint in JOINTS)
    result['MPJAE'] = fmean(errors)
    return result


def index_folder(folder, pattern=FILENAME_PATTERN):
    folder = Path(folder)
    if not folder.is_dir():
        raise FileNotFoundError(f'Folder not found: {folder}')
    regex = re.compile(pattern, re.IGNORECASE)
    if not {'video', 'frame'} <= regex.groupindex.keys():
        raise ValueError('FILENAME_PATTERN needs named video and frame groups')
    indexed = {}
    for path in sorted(folder.rglob('*.json')):
        match = regex.fullmatch(path.stem)
        if not match:
            raise ValueError(f'Cannot parse filename: {path}; adjust FILENAME_PATTERN')
        key = (match['video'].casefold(), int(match['frame']))
        if key in indexed:
            raise ValueError(f'Duplicate pairing key {key}: {indexed[key][0]} and {path}')
        indexed[key] = (path, match['video'])
    if not indexed:
        raise ValueError(f'No JSON files in {folder}')
    return indexed


def load_pairs(pred_folder, gt_folder, pattern=FILENAME_PATTERN, allow_unmatched=False):
    pred = index_folder(pred_folder, pattern)
    gt = index_folder(gt_folder, pattern)
    missing_gt = sorted(pred.keys() - gt.keys())
    missing_pred = sorted(gt.keys() - pred.keys())
    if (missing_gt or missing_pred) and not allow_unmatched:
        raise ValueError(f'Unmatched files: missing GT={missing_gt[:10]}, missing pred={missing_pred[:10]}. '
                         'Fix pairing or explicitly set ALLOW_UNMATCHED=True.')
    keys = sorted(pred.keys() & gt.keys(), key=lambda k: (natural_key(k[0]), k[1]))
    if not keys:
        raise ValueError('No matched frames')
    groups = defaultdict(list)
    manifest = []
    names = {}
    for key in keys:
        pred_path, video = pred[key]
        gt_path, _ = gt[key]
        video = names.setdefault(key[0], video)
        frame = f'{video}_{key[1]:04d}'
        try:
            metrics = calculate_frame(read_points(pred_path), read_points(gt_path))
        except ValueError as exc:
            raise ValueError(f'{frame}: {exc}') from exc
        groups[video].append((frame, metrics))
        manifest.append({'frame': frame, 'pred': str(pred_path.resolve()), 'GT': str(gt_path.resolve())})
    return dict(groups), {'matched_frames': len(keys), 'missing_GT': missing_gt,
                          'missing_pred': missing_pred, 'pairs': manifest}


def summarize(groups):
    rows = []
    for video in sorted(groups, key=natural_key):
        frames = groups[video]
        values = [[record[metric] for _, record in frames] for metric in ('MPJPE', 'MPJAE')]
        rows.append([video, len(frames), *[v for series in values
                                         for v in (fmean(series), min(series), max(series))]])
    # Original has 181 frames per video, so pooled and video-macro means coincide.
    # Explicitly use an equal-weight mean of video means for Overall.
    rows.append(['Overall', None, fmean(row[2] for row in rows), None, None,
                 fmean(row[5] for row in rows), None, None])
    return rows


def write_workbooks(groups, output_folder, overwrite=False):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    output_folder = Path(output_folder)
    detailed = output_folder / 'metrics_per_file_by_video.xlsx'
    summary = output_folder / 'metrics_per_video.xlsx'
    for target in (detailed, summary):
        if target.exists() and not overwrite:
            raise FileExistsError(f'Will not overwrite {target}; change OUTPUT_FOLDER or OVERWRITE')
    for video, frames in groups.items():
        if len(frames) > 16383:
            raise ValueError(f'{video}: too many frames for the original transposed Excel layout')
    output_folder.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    sheet = book.active
    sheet.title = 'metrics_per_file_transposed_by_'
    for i, video in enumerate(sorted(groups, key=natural_key)):
        if i:
            sheet.append([])
        frames = groups[video]
        sheet.append([f'[{video}]'])
        sheet.append(['frame', *[name for name, _ in frames]])
        for metric in frames[0][1]:
            sheet.append([metric, *[record[metric] for _, record in frames]])
    sheet.column_dimensions['A'].width = 27
    sheet.freeze_panes = 'B3'
    sheet.sheet_format.defaultColWidth = 18
    for row in sheet:
        for cell in row:
            if isinstance(cell.value, (float, int)):
                cell.number_format = '0.000000000'
        if row[0].value == 'frame' or str(row[0].value).startswith('['):
            for cell in row:
                if cell.value is not None:
                    cell.font = Font(bold=True)
                    cell.fill = PatternFill('solid', fgColor='DCE6F1')
    book.save(detailed)
    book.close()
    book = Workbook()
    sheet = book.active
    sheet.title = 'metrics_per_video'
    sheet.append(SUMMARY_HEADERS)
    for row in summarize(groups):
        sheet.append(row)
    sheet.freeze_panes = 'C2'
    sheet.auto_filter.ref = f'A1:H{sheet.max_row-1}'
    for column in 'ABCDEFGH':
        sheet.column_dimensions[column].width = 20
    for row in sheet:
        for cell in row:
            if cell.column > 2 and isinstance(cell.value, (int, float)):
                cell.number_format = '0.000000000'
            if cell.row == 1 or cell.row == sheet.max_row:
                cell.font = Font(bold=True)
    book.save(summary)
    book.close()
    return detailed, summary


def run(pred_folder=PRED_FOLDER, gt_folder=GT_FOLDER, output_folder=OUTPUT_FOLDER,
        pattern=FILENAME_PATTERN, allow_unmatched=ALLOW_UNMATCHED, overwrite=OVERWRITE):
    groups, audit = load_pairs(pred_folder, gt_folder, pattern, allow_unmatched)
    detailed, summary = write_workbooks(groups, output_folder, overwrite)
    audit.update(method='12-joint raw-coordinate mean Euclidean error; 8-angle mean absolute error in degrees',
                 alignment='none; no translation, rotation, scaling or temporal alignment',
                 overall='equal-weight mean of video means',
                 detailed=str(detailed.resolve()), summary=str(summary.resolve()))
    (Path(output_folder) / 'pairing_report.json').write_text(
        json.dumps(audit, ensure_ascii=False, indent=2), encoding='utf-8')
    return audit


if __name__ == '__main__':
    report = run()
    print(f"Matched {report['matched_frames']} frames\n{report['detailed']}\n{report['summary']}")
